"""Convert docs/rapport_technique.md to docs/rapport_technique.docx (python-docx).

Supports the subset used by the report: headings, paragraphs, bold/italic/code
inline, bullet and numbered lists (one nesting level), pipe tables and images.
The Mermaid block is replaced by the pre-rendered docs/architecture.png.
Usage: python docs/build_rapport_docx.py
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOCS = Path(__file__).resolve().parent
SRC = DOCS / "rapport_technique.md"
OUT = DOCS / "rapport_technique.docx"
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
BODY_PT, TABLE_PT = 10, 8.5


def add_inline(paragraph, text: str, size: float) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            size_run = size - 0.5
            run.font.size = Pt(size_run)
            continue
        elif part.startswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        run.font.size = Pt(size)


def shade(cell, color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shd)


def add_table(doc, rows: list[list[str]]) -> None:
    header, body = rows[0], [r for r in rows[2:]]
    table = doc.add_table(rows=1 + len(body), cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, values in enumerate([header] + body):
        for c, value in enumerate(values):
            cell = table.cell(r, c)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            add_inline(para, value, TABLE_PT)
            if r == 0:
                shade(cell, "DCE6F1")
                for run in para.runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def main() -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(1.6)
        section.left_margin = section.right_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(BODY_PT)
    normal.paragraph_format.space_after = Pt(3)
    for level, size in ((1, 15), (2, 12.5)):
        style = doc.styles[f"Heading {level}"]
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        style.paragraph_format.space_before = Pt(8 if level == 2 else 0)
        style.paragraph_format.space_after = Pt(3)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            fence_end = next(j for j in range(i + 1, len(lines)) if lines[j].startswith("```"))
            if not line.startswith("```mermaid"):
                para = doc.add_paragraph()
                add_inline(para, "`" + "\n".join(lines[i + 1:fence_end]) + "`", BODY_PT)
            i = fence_end + 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue
        image = re.match(r"!\[[^\]]*\]\(([^)]+)\)", line)
        if image:
            doc.add_picture(str(DOCS / image.group(1)), height=Cm(10.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption = doc.add_paragraph()
            add_inline(caption, "*Figure 1 — Architecture multi-agents (source Mermaid dans rapport_technique.md).*", 8.5)
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif re.match(r"^\s*- ", line):
            nested = line.startswith("   ")
            para = doc.add_paragraph(style="List Bullet 2" if nested else "List Bullet")
            para.paragraph_format.space_after = Pt(1)
            add_inline(para, line.strip()[2:], BODY_PT)
        elif re.match(r"^\d+\. ", line):
            para = doc.add_paragraph(style="List Number")
            para.paragraph_format.space_after = Pt(1)
            add_inline(para, re.sub(r"^\d+\. ", "", line), BODY_PT)
        elif line.strip():
            para = doc.add_paragraph()
            add_inline(para, line, BODY_PT)
        i += 1
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
